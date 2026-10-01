import sqlite3
import tempfile
import unittest
from pathlib import Path
from threading import Event
from openpyxl import load_workbook
from repositories.machine_slot_summary_repository import load_machine_slot_summary
from services.machine_slot_export import export_machine_slots, slot_metrics


class ExportTests(unittest.TestCase):
    def test_range_empty_slots_history_and_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            db, out = Path(folder)/'data.db', Path(folder)/'result.xlsx'
            with sqlite3.connect(db) as c:
                c.execute('''CREATE TABLE prime_data(id INTEGER PRIMARY KEY, EQP TEXT, SLOT INTEGER,
                    DATE TEXT,TIME TEXT,RESULT TEXT,TEST_COUNT INTEGER,QTY INTEGER,MODEL TEXT,
                    SCRAPCODE TEXT,PARTNO TEXT,LOTNO TEXT,SERIAL TEXT,TIER TEXT)''')
                c.execute('CREATE TABLE slot_fail_alarm(eqp TEXT,slot INTEGER,alarm_date TEXT)')
                c.execute('''CREATE TABLE machine_slot_yield_config(id INTEGER,target_15 REAL,
                    target_30 REAL,continuous_fail_count INTEGER,different_scrap_fail_count INTEGER)''')
                c.execute('INSERT INTO machine_slot_yield_config VALUES(1,50,70,3,4)')
                c.executemany('INSERT INTO slot_fail_alarm VALUES(?,?,?)',
                              [('A',1,'20260901'),('A',2,'20260920')])
                for slot in (1,2):
                    for i, result in enumerate(['FAIL','FAIL','FAIL','PASS']):
                        c.execute('''INSERT INTO prime_data(EQP,SLOT,DATE,TIME,RESULT,TEST_COUNT,QTY,
                            MODEL,SCRAPCODE,PARTNO,LOTNO,SERIAL,TIER) VALUES(?,?,?,?,?,0,1,'MODEL','E1','PART','LOT',?,'')''',
                            ('A',slot,'20260920',f'10:00:0{i}',result,f'=ID{i}'))
                c.execute("INSERT INTO prime_data(EQP,SLOT,DATE,TIME,RESULT,TEST_COUNT,QTY) VALUES('B',1,'20260801','10:00:00','PASS',0,1)")
            payload = load_machine_slot_summary(db, '20260920', '20260921', Event())
            self.assertEqual(len(payload['rows']),2)
            self.assertEqual([r['alarm'] for r in payload['rows']],[False,True])
            self.assertEqual(len(payload['daily']),192)  # zero rows stay in Daily
            export_machine_slots(db,out,('20260920','20260921'),payload['rows'],payload['daily'],Event())
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames,['slot summary','Daily','test history'])
            s=wb.worksheets[0]
            self.assertEqual((s['R2'].value,s['T2'].value),(.25,.25))
            self.assertEqual((s['S2'].value,s['U2'].value),(50,70))
            self.assertEqual((s['V2'].value,s['W2'].value),('YES','YES'))
            self.assertEqual((s['Z2'].value,s['Z3'].value),('PASS','ALARM'))
            self.assertEqual(s['A3'].fill.fgColor.rgb,'00FCA5A5')
            self.assertEqual(wb.worksheets[1].max_row,193)
            h=wb.worksheets[2]
            self.assertEqual(h.max_row,9)
            self.assertEqual(h['K2'].data_type,'s')
            self.assertEqual(h['K2'].value,'=ID3')
            wb.close()

    def test_latest_windows_and_retest_exclusion(self):
        config=dict(target_15=50,target_30=70,continuous_fail_count=3,different_scrap_fail_count=4)
        events=[dict(RESULT='PASS' if i<15 else 'FAIL',TEST_COUNT=0,MODEL='M',SCRAPCODE='E') for i in range(31)]
        self.assertEqual(slot_metrics(events,config)[:2],(1,.5))
        events=[dict(RESULT='FAIL',TEST_COUNT=0,MODEL='M',SCRAPCODE=s) for s in ['A','B','C','D']]
        self.assertEqual(slot_metrics(events,config)[2:],(False,True))
        events[2]['TEST_COUNT']=1
        self.assertEqual(slot_metrics(events,config)[2:],(False,False))

if __name__ == '__main__':
    unittest.main()
