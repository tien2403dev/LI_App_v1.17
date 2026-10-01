import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import csv
import io
import tempfile
import unittest
from pathlib import Path
from dataclasses import replace
from PyQt5.QtWidgets import QApplication
from openpyxl import load_workbook
from repositories.report_repository import ReportSummaryRow
from ui.widgets.report_builder import ReportBuilder
from services.report_export import clipboard_text, export_excel


class BuilderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.first = ReportSummaryRow("20260909", "AI-H911", 47, 79, 79, 0, "", 100, "MZWL6")
        self.second = ReportSummaryRow("20260909", "AI-H911", 46, 83, 78, 5,
                                       "4541*4\n4593*1", 100*78/83, "MZWL6")
        self.widget = ReportBuilder()

    def tearDown(self):
        self.widget.close()

    def test_add_merge_deduplicate_and_update(self):
        w = self.widget
        w.add_rows([self.first, self.second])
        self.assertEqual([r.slot for r in w.rows], [47,46])
        self.assertEqual(w.table.rowSpan(0,0), 2)
        self.assertEqual(w.table.rowSpan(0,1), 2)
        w.add_rows([replace(self.first, in_qty=80)])
        self.assertEqual(len(w.rows),2)
        self.assertEqual(w.rows[0].in_qty,80)

    def test_move_and_remove(self):
        w = self.widget
        w.add_rows([self.first,self.second])
        group = w.tree.topLevelItem(0).child(0)
        w.tree.setCurrentItem(group.child(1))
        w.move(-1)
        self.assertEqual([r.slot for r in w.rows],[46,47])
        w.tree.setCurrentItem(group)
        w.remove_selected()
        self.assertEqual(w.rows,[])
        self.assertEqual(w.entries,{})
        self.assertEqual(w.tree.topLevelItemCount(),0)

    def test_copy_zero_and_no_extra_lines(self):
        rows=list(csv.reader(io.StringIO(clipboard_text([self.first,self.second])),delimiter="\t"))
        self.assertEqual(len(rows),3)
        self.assertEqual(rows[1][5],"0")
        self.assertEqual(rows[2][6],"4541*4 / 4593*1")

    def test_excel_merges_numbers_colors_and_blanks(self):
        empty=replace(self.first,date="20260910",in_qty=None,out_qty=None,fail_qty=None,yield_percent=None)
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"out.xlsx"
            export_excel(path,[self.first,self.second,empty])
            wb=load_workbook(path)
            ws=wb.active
            self.assertIn("A2:A3",str(ws.merged_cells))
            self.assertIn("B2:B3",str(ws.merged_cells))
            self.assertEqual(ws["F2"].value,0)
            self.assertEqual(ws["H2"].value,1)
            self.assertEqual(ws["H2"].number_format,"0.00%")
            self.assertEqual(ws["G3"].value,"4541*4\n4593*1")
            self.assertEqual(ws["H2"].fill.fgColor.rgb,"0063BE7B")
            self.assertIsNone(ws["D4"].value)
            wb.close()


if __name__ == "__main__":
    unittest.main()
