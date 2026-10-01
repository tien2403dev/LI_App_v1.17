import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import sqlite3, tempfile, time, unittest
from pathlib import Path
from PyQt5.QtWidgets import QApplication
from PyQt5.QtCore import QDateTime
from openpyxl import load_workbook
from services.slot_report_service import load_groups,refresh_groups,merged_spans,export_excel,clipboard_html
from ui.pages.report_page import ReportPage
from ui.widgets.slot_report_builder import SlotPicker

class SlotReportTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls): cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory(); self.db=Path(self.tmp.name)/'test.db'
  with sqlite3.connect(self.db) as c:
   c.execute('CREATE TABLE prime_data (id INTEGER PRIMARY KEY, DATE TEXT,TIME TEXT,EQP TEXT,Slot INTEGER,RESULT TEXT,SCRAPCODE TEXT,MODEL TEXT,TEST_COUNT INTEGER,PARTNO TEXT,LOTNO TEXT,SERIAL TEXT,QTY INTEGER,TIER TEXT)')
   for hour,slot,result,code,count in [('08:00:00',47,'PASS','',0),('08:10:00',46,'FAIL','4541',0),('08:20:00',46,'FAIL','4593',0),('08:30:00',46,'PASS','',1),('09:00:00',47,'FAIL','X',0),('09:10:00',46,'PASS','',0)]:
    c.execute('INSERT INTO prime_data (DATE,TIME,EQP,Slot,RESULT,SCRAPCODE,MODEL,TEST_COUNT,QTY) VALUES (?,?,?,?,?,?,?,?,?)',('20260909',hour,'AI-H911',slot,result,code,'M',count,1))
 def tearDown(self): self.tmp.cleanup()
 def groups(self,first='08:00:00',last='08:59:59',slots=(47,46,48)):
  return load_groups(self.db,'2026-09-09 '+first,'2026-09-09 '+last,'AI-H911',slots)
 def wait(self,predicate):
  deadline=time.monotonic()+5
  while not predicate() and time.monotonic()<deadline:
   self.app.processEvents(); time.sleep(.005)
  self.assertTrue(predicate())
 def test_count_order_empty_time(self):
  rows=self.groups()[0].rows
  self.assertEqual([r.slot for r in rows],[47,46,48])
  self.assertEqual([(r.in_qty,r.out_qty,r.fail_qty) for r in rows],[(1,1,0),(2,0,2),(None,None,None)])
  self.assertEqual(set(rows[1].scrap_codes.splitlines()),{'4541*1','4593*1'})
  self.assertEqual(self.groups('09:00:00','09:59:59')[0].rows[0].fail_qty,1)
 def test_refresh_merge_export(self):
  groups=self.groups()+self.groups('09:00:00','09:59:59',(47,46))
  self.assertEqual(refresh_groups(self.db,groups),groups)
  self.assertEqual(list(merged_spans(groups)),[(0,3,0),(0,3,1),(3,2,0),(3,2,1)])
  path=Path(self.tmp.name)/'report.xlsx'; export_excel(path,groups)
  wb=load_workbook(path); ws=wb.active
  self.assertEqual(ws.max_column,9)
  self.assertEqual(set(map(str,ws.merged_cells.ranges)),{'A2:A4','B2:B4','A5:A6','B5:B6'})
  self.assertEqual(ws['F2'].value,0); self.assertIsNone(ws['D4'].value)
  self.assertIn('\n',ws['G3'].value); self.assertEqual(ws['H2'].value,1)
  self.assertEqual(ws['G3'].font.color.rgb,'00FF0000')
  self.assertIn('rowspan="3"',clipboard_html(groups)); self.assertIn('<br>',clipboard_html(groups)); wb.close()
 def test_picker(self):
  p=SlotPicker([47,46]); self.assertEqual(len(p.boxes),48)
  grid=p.boxes[1].parentWidget().layout()
  for slot in (1,12,13,24,25,36,37,48):
   self.assertEqual(grid.getItemPosition(grid.indexOf(p.boxes[slot]))[:2],((slot-1)%12+1,(slot-1)//12))
  p.boxes[1].click(); self.assertEqual(p.order,[47,46,1]); p.close()
 def test_gui_flow(self):
  page=ReportPage(self.db); page.resize(1450,700); page.show()
  try:
   b=page.builder; self.wait(lambda:not page.has_running_tasks())
   self.assertEqual([page.pages.tabText(i) for i in range(2)],['Báo cáo Slot','Tra cứu SLOT'])
   b.first.setDateTime(QDateTime.fromString('2026-09-09 08:00:00','yyyy-MM-dd HH:mm:ss'))
   b.last.setDateTime(QDateTime.fromString('2026-09-09 08:59:59','yyyy-MM-dd HH:mm:ss'))
   b.eqp.setCurrentText('AI-H911'); b.slots=[47,46,48]; b._selection_changed()
   self.wait(lambda:b.add_button.isEnabled()); self.assertEqual(b.preview.rowSpan(0,1),3)
   b._add(); b._add(); self.assertEqual(len(b.groups),1)
   b.last.setDateTime(QDateTime.fromString('2026-09-09 09:59:59','yyyy-MM-dd HH:mm:ss'))
   self.assertFalse(b.add_button.isEnabled()); self.wait(lambda:b.add_button.isEnabled()); b._add()
   self.assertEqual(len(b.groups),2); b._refresh(); self.wait(lambda:b.thread is None)
   b._details(b.groups,4); self.wait(lambda:not b.is_busy())
   self.assertEqual(b.dialogs[0].model.rowCount(),3)
   self.assertTrue(all(r.test_count==0 for r in b.dialogs[0].model.rows)); b.dialogs[0].close()
   b.order_list.setCurrentRow(1); b._move(-1); b._remove(); self.assertEqual(len(b.groups),1)
   page.grab().save('/workspace/scratch/875308bb2752/report_check.png')
  finally:
   page.prepare_close(); self.wait(lambda:not page.has_running_tasks()); page.close()
 def test_no_creation(self):
  missing=Path(self.tmp.name)/'missing.db'
  with self.assertRaises(sqlite3.OperationalError):
   load_groups(missing,'2026-09-09 00:00:00','2026-09-09 23:59:59','AI-H911',[1])
  self.assertFalse(missing.exists())
