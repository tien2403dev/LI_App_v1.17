"""Slot report composition/export; counts all test attempts."""
from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from html import escape
from pathlib import Path
import csv
import io
import os
import tempfile

from repositories.report_repository import ReportRepository

HEADERS = ['Date', 'EQP', 'SLOT', 'IN', 'OutQty', 'FailQty', 'Scrap Code', 'YLD', 'Note']
WIDTHS = [14, 15, 9, 9, 10, 10, 40, 12, 28]


class SlotRepository(ReportRepository):
    def get_eqps(self):
        """List every machine present in PRIME; report rows include all TEST_COUNT values."""
        with self._connect() as connection:
            # Same index-seek strategy as the main filter: visit each EQP
            # once instead of scanning every test and sorting TRIM(EQP).
            rows = connection.execute("""WITH RECURSIVE options(value) AS (
                SELECT MIN(EQP) FROM prime_data
                UNION ALL
                SELECT (SELECT MIN(EQP) FROM prime_data WHERE EQP>options.value)
                FROM options WHERE value IS NOT NULL
            ) SELECT value FROM options WHERE value IS NOT NULL""")
            # Preserve SQLite TRIM's space-only normalization and deduplication.
            return sorted({row[0].strip(' ') for row in rows if row[0].strip(' ')})


@dataclass(frozen=True)
class SlotGroup:
    date: str
    eqp: str
    slots: tuple
    rows: tuple
    date_from: str
    date_to: str

    @property
    def key(self):
        return self.date, self.eqp, self.slots, self.date_from, self.date_to


def load_groups(database_path, first, last, eqp, slots=()):
    datetime.strptime(first, '%Y-%m-%d %H:%M:%S')
    datetime.strptime(last, '%Y-%m-%d %H:%M:%S')
    if first > last:
        raise ValueError('From phải nhỏ hơn hoặc bằng To.')
    slots = tuple(dict.fromkeys(slots))
    if any(type(s) is not int or not 1 <= s <= 48 for s in slots):
        raise ValueError('Slot phải từ 1 đến 48.')
    if not eqp or not slots:
        return []
    rows = SlotRepository(database_path).get_summary(first, last, [eqp], list(slots))
    dates = OrderedDict()
    for row in rows:
        dates.setdefault(row.date, []).append(row)
    groups = []
    for date, items in dates.items():
        day = datetime.strptime(date, '%Y%m%d').strftime('%Y-%m-%d')
        groups.append(SlotGroup(date, eqp, slots, tuple(items),
                               max(first, day + ' 00:00:00'), min(last, day + ' 23:59:59')))
    return groups


def refresh_groups(database_path, groups):
    return [load_groups(database_path, g.date_from, g.date_to, g.eqp, g.slots)[0]
            for g in groups]


def report_rows(groups):
    """Flatten groups to typed cell values shared by preview/clipboard/Excel."""
    def slash_text(value):
        """Keep multiple database values in one cell, separated by ' / '."""
        normalized = str(value or '').replace('\r\n', '\n').replace('\r', '\n')
        return ' / '.join(part.strip() for part in normalized.split('\n') if part.strip())

    result = []
    for group in groups:
        date = datetime.strptime(group.date, '%Y%m%d')
        for row in group.rows:
            result.append([date, group.eqp, row.slot, row.in_qty, row.out_qty, row.fail_qty,
                           (slash_text(row.scrap_codes) or '-') if row.in_qty is not None else '',
                           None if row.yield_percent is None else row.yield_percent / 100,
                           slash_text(row.models) if row.in_qty is not None else ''])
    return result


def display_value(value, column):
    """Format only at presentation boundaries, preserving zero quantities."""
    if value is None:
        return ''
    if column == 0:
        return value.strftime('%d/%m/%Y')
    if column == 7:
        return f'{value:.2%}'
    return str(value)


def merged_spans(groups):
    """Each Add selection/day is its own Date/EQP merge block."""
    offset = 0
    for group in groups:
        count = len(group.rows)
        if count > 1:
            for column in (0, 1):
                yield offset, count, column
        offset += count


def yield_color(value):
    """Sample report colors: >=95% yellow, otherwise orange."""
    if value is None:
        return None
    return 'FFEB84' if value >= .95 else 'F7A16C'


def clipboard_text(groups):
    """Plain TSV repeats group labels, keeping one physical line per slot."""
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, delimiter='\t', lineterminator='\r\n')
    for row in report_rows(groups):
        values = [display_value(v, c).replace('\n', ' / ').replace('\r', '')
                  for c, v in enumerate(row)]
        # Prevent text labels from being interpreted as spreadsheet formulas.
        for c in (1, 6, 8):
            if values[c] != '-' and values[c].startswith(('=', '+', '-', '@')):
                values[c] = "'" + values[c]
        writer.writerow(values)
    return stream.getvalue()


def clipboard_html(groups):
    """HTML with cell merges and colors for rich clipboard consumers."""
    spans = {(r, c): count for r, count, c in merged_spans(groups)}
    covered = {(r + offset, c) for (r, c), count in spans.items()
               for offset in range(1, count)}
    parts = ['<html><body><table style="border-collapse:collapse;text-align:center;">']
    for r, row in enumerate(report_rows(groups)):
        parts.append('<tr>')
        for c, value in enumerate(row):
            if (r, c) in covered:
                continue
            style = 'border:1px solid black;padding:5px;vertical-align:middle;white-space:pre-wrap;'
            if c in (1, 6, 8):
                style += 'mso-number-format:"\\@";'
            color = yield_color(value) if c == 7 else None
            if c == 6:
                style += 'color:#FF0000;'
            if color:
                style += f'background:#{color};'
            parts.append(f'<td rowspan="{spans.get((r, c), 1)}" style="{escape(style, quote=True)}">'
                         + escape(display_value(value, c)).replace('\n', '<br>') + '</td>')
        parts.append('</tr>')
    return ''.join(parts) + '</table></body></html>'


def export_excel(path, groups):
    """Write formatted report atomically, preserving an existing file on failure."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Slot Report'
    ws.append(HEADERS)
    thin = Side(style='thin', color='000000')
    for r, values in enumerate(report_rows(groups), 2):
        ws.append(values)
        for cell in ws[r]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
        ws.cell(r, 1).number_format = 'dd/mm/yyyy'
        ws.cell(r, 8).number_format = '0.00%'
        color = yield_color(values[7])
        if color:
            ws.cell(r, 8).fill = PatternFill('solid', fgColor=color)
        ws.row_dimensions[r].height = max(26, 16 * max(
            len(str(values[c] or '').splitlines()) for c in (6, 8)))
    for row in ws:
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7).font = Font(name='Calibri', size=11, color='FF0000')
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='B7DEE8')
        cell.font = Font(name='Calibri', bold=True, size=11)
    ws.row_dimensions[1].height = 28
    for start, count, col in merged_spans(groups):
        ws.merge_cells(start_row=start + 2, end_row=start + count + 1,
                       start_column=col + 1, end_column=col + 1)
    for c, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = 'D2'
    ws.print_title_rows = '1:1'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    path = Path(path)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as f:
            temporary = f.name
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        wb.close()
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
