"""Presentation/export only; all quantities come from ReportRepository."""
from datetime import datetime
import csv
import io

HEADERS = ["Date", "EQP", "SLOT", "IN", "OutQty", "FailQty", "Scrap Code", "YLD"]


def cells(row):
    date = datetime.strptime(row.date, "%Y%m%d").strftime("%d/%m/%Y")
    if row.in_qty is None:
        return [date, row.eqp, row.slot, "", "", "", "", ""]
    return [date, row.eqp, row.slot, row.in_qty, row.out_qty, row.fail_qty,
            row.scrap_codes.replace(" / ", "\n") or "-",
            "" if row.yield_percent is None else row.yield_percent / 100]


def yield_color(percent):
    if percent is None:
        return None
    return "63BE7B" if percent >= 100 else "FFEB84" if percent >= 95 else "F7A16C"


def spans(rows):
    """Only merge adjacent date/EQP groups, preserving user order."""
    for column, key in ((0, lambda r: r.date), (1, lambda r: (r.date, r.eqp))):
        start = 0
        while start < len(rows):
            end = start + 1
            while end < len(rows) and key(rows[end]) == key(rows[start]):
                end += 1
            if end - start > 1:
                yield start, end - start, column
            start = end


def clipboard_text(rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter="\t", lineterminator="\r\n")
    writer.writerow(HEADERS)
    for row in rows:
        values = cells(row)
        values[6] = values[6].replace("\n", " / ")
        if isinstance(values[7], (int, float)):
            values[7] = f"{values[7]:.2%}"
        writer.writerow(values)
    return stream.getvalue()


def export_excel(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(HEADERS)
    border = Border(*( [Side(style="thin", color="000000")] * 4 ))
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="B7DEE8")
        cell.font = Font(bold=True, size=11)
    for index, row in enumerate(rows, 2):
        ws.append(cells(row))
        # Database text is always text, never an Excel formula.
        for cell in ws[index]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        ws.cell(index, 7).font = Font(color="FF0000", size=11)
        ws.cell(index, 8).number_format = "0.00%"
        color = yield_color(row.yield_percent) if row.in_qty is not None else None
        if color:
            ws.cell(index, 8).fill = PatternFill("solid", fgColor=color)
        ws.row_dimensions[index].height = max(25, 18 * len(cells(row)[6].split("\n")))
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    for start, count, column in spans(rows):
        ws.merge_cells(start_row=start+2, end_row=start+count+1,
                       start_column=column+1, end_column=column+1)
    for i, width in enumerate((15, 16, 10, 11, 11, 11, 24, 13), 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    wb.save(path)
