"""Background Excel export of the applied Machine Slot views (read-only DB)."""
from collections import deque
from pathlib import Path
from tempfile import NamedTemporaryFile
import os

from database.read_snapshot import machine_export_snapshot

SUMMARY_HEADERS = ['No', 'Machine', 'Slot', 'Total Test', 'PASS', 'FAIL', 'YIELD', 'Fail comment'] + [
    f'Lần {i}' for i in range(1, 11)] + ['Last 15 Yield', 'Target 15',
    'Last 30 Yield', 'Target 30', 'Fail Consecutive', 'Same Scrapcode Fail',
    'Alarm 15', 'Alarm 30', 'ALARM']
DAILY_HEADERS = ['No', 'Date', 'Machine', 'Slot', 'Total Test', 'Pass', 'Fail', 'Yield']
HISTORY_HEADERS = ['DATE', 'TIME', 'PARTNO', 'LOTNO', 'SERIAL', 'RESULT',
                   'SCRAPCODE', 'QTY', 'EQP', 'TEST_COUNT', 'Slot', 'MODEL', 'TIER']


def check_cancel(cancel):
    if cancel.is_set():
        raise InterruptedError('Đã hủy xuất Excel.')


def slot_metrics(events, config):
    """Newest-first rows; only consecutive rules ignore TEST_COUNT != 0."""
    rates = []
    for size in (15, 30):
        sample = events[:size]
        passed = sum(r['RESULT'] == 'PASS' for r in sample)
        failed = sum(r['RESULT'] == 'FAIL' for r in sample)
        rates.append(passed / (passed + failed) if passed + failed else 0)
    same = mixed = False
    key, count = None, 0
    queue = deque(maxlen=int(config['different_scrap_fail_count']))
    for r in events:
        if r['TEST_COUNT'] != 0:
            continue
        current = (r['MODEL'], r['SCRAPCODE']) if r['RESULT'] == 'FAIL' and r['SCRAPCODE'] else None
        count = count + 1 if current is not None and current == key else (1 if current else 0)
        key = current
        same |= count >= int(config['continuous_fail_count'])
        if current is None or (queue and queue[-1]['MODEL'] != r['MODEL']):
            queue.clear()
        if current:
            queue.append(r)
            mixed |= len(queue) == queue.maxlen and len({e['SCRAPCODE'] for e in queue}) > 1
    return rates[0], rates[1], same, mixed


def export_machine_slots(database_path, output, dates, rows, daily, cancel):
    # Lazy dependency: ordinary table loading does not import openpyxl.
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(write_only=True)
    summary = wb.create_sheet('slot summary')
    day_sheet = wb.create_sheet('Daily')
    history = wb.create_sheet('test history')
    counts = {}

    def append(sheet, values, *, header=False, alarm=False, percentages=()):
        check_cancel(cancel)
        counts[sheet.title] = counts.get(sheet.title, 0) + 1
        if counts[sheet.title] > 1048576:
            raise ValueError('Vượt giới hạn dòng Excel. Hãy chọn khoảng ngày hoặc số slot ít hơn.')
        cells = []
        for col, value in enumerate(values, 1):
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'  # Preserve IDs/text; never execute log text as formulas.
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, bold=header,
                             color='FFFFFF' if header else '172B4D')
            if header or alarm:
                cell.fill = PatternFill('solid', fgColor='214666' if header else 'FCA5A5')
            if col in percentages and not header:
                cell.number_format = '0.00%'
            if not header and value == 'FAIL':
                cell.font = Font(name='Calibri', size=11, color='DC2626')
            elif not header and value == 'PASS':
                cell.font = Font(name='Calibri', size=11, color='16845B')
            cells.append(cell)
        sheet.append(cells)

    for sheet, headers in ((summary, SUMMARY_HEADERS), (day_sheet, DAILY_HEADERS), (history, HISTORY_HEADERS)):
        sheet.freeze_panes = 'D2'
        sheet.row_dimensions[1].height = 34
        for i, label in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = max(12, min(24, len(label) + 3))
        if sheet is history:
            sheet.column_dimensions['C'].width = 28
            sheet.column_dimensions['D'].width = 20
            sheet.column_dimensions['E'].width = 23
        append(sheet, headers, header=True)
    temp = None
    try:
        with machine_export_snapshot(database_path, dates, cancel) as conn:
            conn.execute('PRAGMA query_only=ON')
            conn.set_progress_handler(lambda: int(cancel.is_set()), 1000)
            conn.execute('BEGIN')
            config_row = conn.execute('SELECT * FROM machine_slot_yield_config WHERE id=1').fetchone()
            if config_row is None:
                raise ValueError('Chưa lưu cấu hình Machine Slot Yield.')
            config = dict(config_row)
            alarms = {(r['eqp'], r['slot']) for r in conn.execute(
                'SELECT eqp,slot FROM slot_fail_alarm WHERE alarm_date BETWEEN ? AND ?', dates)}
            for number, row in enumerate(rows, 1):
                check_cancel(cancel)
                events = [dict(r) for r in conn.execute('''SELECT * FROM prime_data
                    WHERE EQP=? AND SLOT=? AND DATE BETWEEN ? AND ?
                    ORDER BY DATE DESC,TIME DESC,id DESC''', (row['eqp'], row['slot'], *dates))]
                if (len(events), sum(r['RESULT']=='PASS' for r in events), sum(r['RESULT']=='FAIL' for r in events)) != (row['total'], row['passed'], row['failed']):
                    raise ValueError('Dữ liệu vừa thay đổi. Bấm Search rồi xuất Excel lại.')
                y15, y30, same, mixed = slot_metrics(events, config)
                recent = [f"{r['RESULT']}\n{r['DATE'][:4]}-{r['DATE'][4:6]}-{r['DATE'][6:]}" for r in events[:10]]
                recent += [''] * (10-len(recent))
                alarm = (row['eqp'], row['slot']) in alarms
                values = [number, row['eqp'], row['slot'], row['total'], row['passed'], row['failed'],
                          row['passed']/row['total'], row.get('fail_comment', ''), *recent, y15, config['target_15'],
                          y30, config['target_30'], 'YES' if same or mixed else 'NO',
                          'YES' if same else 'NO', 'YES' if y15*100 < config['target_15'] else 'NO',
                          'YES' if y30*100 < config['target_30'] else 'NO', 'ALARM' if alarm else 'PASS']
                summary.row_dimensions[number+1].height = 34
                append(summary, values, alarm=alarm, percentages=(7,19,21))
                # Finish all dates for this EQP/Slot before advancing to the next.
                for r in events:
                    append(history, [r['DATE'], r['TIME'], r['PARTNO'], r['LOTNO'],
                        r['SERIAL'], r['RESULT'], r['SCRAPCODE'], r['QTY'], r['EQP'],
                        r['TEST_COUNT'], r['SLOT'], r['MODEL'], r.get('TIER', '')])
            for i, r in enumerate(daily, 1):
                append(day_sheet, [i, r['date'], r['eqp'], r['slot'], r['total'], r['passed'],
                    r['failed'], r['passed']/r['total'] if r['total'] else None], percentages=(8,))
        for sheet in wb:
            sheet.auto_filter.ref = f'A1:{get_column_letter(len(SUMMARY_HEADERS) if sheet is summary else len(DAILY_HEADERS) if sheet is day_sheet else len(HISTORY_HEADERS))}{counts[sheet.title]}'
        with NamedTemporaryFile(dir=Path(output).resolve().parent, suffix='.xlsx', delete=False) as f:
            temp = f.name
        wb.save(temp)
        check_cancel(cancel)
        os.replace(temp, output)
    finally:
        for sheet in wb:
            if not sheet.closed:
                sheet.close()
        wb.close()
        if temp and os.path.exists(temp):
            os.unlink(temp)
